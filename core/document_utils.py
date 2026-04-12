"""File and document text extraction helpers with graceful optional dependency handling."""

from __future__ import annotations

import csv
import json
import zipfile
from pathlib import Path
from typing import Any, Dict, Tuple


def _extract_text_plain(path: Path) -> Tuple[str, str]:
    content = path.read_text(encoding="utf-8", errors="ignore")
    return content, "plain_text"


def _extract_text_csv(path: Path) -> Tuple[str, str]:
    rows = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
        reader = csv.reader(handle)
        for row in reader:
            rows.append(", ".join(str(cell) for cell in row))
    return "\n".join(rows), "csv"


def _extract_text_json(path: Path) -> Tuple[str, str]:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return json.dumps(data, indent=2, ensure_ascii=False), "json"


def _extract_text_html(path: Path) -> Tuple[str, str]:
    import re

    raw = path.read_text(encoding="utf-8", errors="ignore")
    text = re.sub(r"<[^>]+>", " ", raw)
    return re.sub(r"\s+", " ", text).strip(), "html_regex"


def _extract_text_pdf(path: Path) -> Tuple[str, str]:
    try:
        import pypdf  # type: ignore

        reader = pypdf.PdfReader(str(path))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages), "pypdf"
    except Exception as exc:
        return "", f"pdf_unavailable:{exc}"


def _extract_text_docx(path: Path) -> Tuple[str, str]:
    try:
        import docx  # type: ignore

        document = docx.Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return text, "python-docx"
    except Exception as exc:
        return "", f"docx_unavailable:{exc}"


def _extract_text_xlsx(path: Path) -> Tuple[str, str]:
    try:
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        lines = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) for cell in row if cell is not None]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines), "openpyxl"
    except Exception as exc:
        return "", f"xlsx_unavailable:{exc}"


def _extract_text_pptx(path: Path) -> Tuple[str, str]:
    try:
        import pptx  # type: ignore

        presentation = pptx.Presentation(str(path))
        chunks = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                text = getattr(shape, "text", "")
                if text:
                    chunks.append(text)
        return "\n".join(chunks), "python-pptx"
    except Exception as exc:
        return "", f"pptx_unavailable:{exc}"


def _extract_text_zip_xml(path: Path, xml_name_hint: str) -> Tuple[str, str]:
    try:
        import re

        with zipfile.ZipFile(path, "r") as archive:
            candidates = [name for name in archive.namelist() if xml_name_hint in name and name.endswith(".xml")]
            texts = []
            for name in candidates:
                raw = archive.read(name).decode("utf-8", errors="ignore")
                cleaned = re.sub(r"<[^>]+>", " ", raw)
                texts.append(re.sub(r"\s+", " ", cleaned).strip())
            return "\n".join(texts), "zip_xml_fallback"
    except Exception as exc:
        return "", f"zip_xml_unavailable:{exc}"


def extract_text_from_file(path_str: str) -> Dict[str, Any]:
    """Extract text and extraction metadata from a local file path."""
    path = Path(path_str)
    if not path.exists() or not path.is_file():
        return {
            "status": "error",
            "error": "file does not exist",
            "path": str(path),
            "text": "",
            "extractor": None,
        }

    suffix = path.suffix.lower()
    extractor = ""
    text = ""

    try:
        if suffix in {".txt", ".log", ".md", ".rst"}:
            text, extractor = _extract_text_plain(path)
        elif suffix == ".csv":
            text, extractor = _extract_text_csv(path)
        elif suffix == ".json":
            text, extractor = _extract_text_json(path)
        elif suffix in {".html", ".htm"}:
            text, extractor = _extract_text_html(path)
        elif suffix == ".pdf":
            text, extractor = _extract_text_pdf(path)
        elif suffix == ".docx":
            text, extractor = _extract_text_docx(path)
            if not text:
                text, extractor = _extract_text_zip_xml(path, "word/")
        elif suffix == ".xlsx":
            text, extractor = _extract_text_xlsx(path)
            if not text:
                text, extractor = _extract_text_zip_xml(path, "xl/")
        elif suffix == ".pptx":
            text, extractor = _extract_text_pptx(path)
            if not text:
                text, extractor = _extract_text_zip_xml(path, "ppt/")
        else:
            text, extractor = _extract_text_plain(path)
    except Exception as exc:
        return {
            "status": "error",
            "error": str(exc),
            "path": str(path),
            "text": "",
            "extractor": None,
        }

    return {
        "status": "ok",
        "path": str(path),
        "text": text,
        "extractor": extractor,
        "char_count": len(text),
    }
